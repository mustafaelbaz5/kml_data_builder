"""KML/KMZ parsing, validation, repair, and export for land parcel review."""

from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass, field
from html import unescape
from html.parser import HTMLParser
from pathlib import Path
import copy
import io
import re
import uuid
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

KML_NS = "http://www.opengis.net/kml/2.2"
GX_NS = "http://www.google.com/kml/ext/2.2"
NS = {"k": KML_NS}
ET.register_namespace("", KML_NS)
ET.register_namespace("gx", GX_NS)

UUID_RE = re.compile(
    r"^[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[1-5][0-9a-fA-F]{3}-[89abAB][0-9a-fA-F]{3}-[0-9a-fA-F]{12}$"
)
ANY_UUID_RE = re.compile(
    r"\b[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[1-5][0-9a-fA-F]{3}-[89abAB][0-9a-fA-F]{3}-[0-9a-fA-F]{12}\b"
)

FIELDS = [
    ("id", "كود القطعة"),
    ("holding_number", "رقم الحيازة"),
    ("parcel_count", "عدد القطع بالحيازة"),
    ("holder_name", "اسم الحائز"),
    ("owner_name", "اسم المالك"),
    ("national_id", "الرقم القومي للحائز"),
    ("association_name", "اسم الجمعية"),
    ("association_code", "كود الجمعية"),
    ("basin_name", "اسم الحوض"),
    ("basin_code", "كود الحوض"),
    ("parcel_number", "رقم القطعة"),
    ("feddan", "فدان"),
    ("qirat", "قيراط"),
    ("sahm", "سهم"),
    ("land_use", "نوع الاستخدام"),
    ("crop_type", "نوع المحصول"),
    ("growth_stage", "مراحل النمو"),
    ("notes", "ملاحظات علي الحيازة (من فريق العمل)"),
]

LABELS = {
    "id": ["id", "ID", "كود القطعة"],
    "holding_number": ["رقم الحيازة"],
    "parcel_count": ["عدد القطع بالحيازة", "عدد القطع"],
    "holder_name": ["اسم الحائز"],
    "owner_name": ["اسم المالك"],
    "national_id": ["الرقم القومي للحائز", "الرقم القومي"],
    "association_name": ["اسم الجمعية"],
    "association_code": ["كود الجمعية"],
    "basin_name": ["اسم الحوض"],
    "basin_code": ["كود الحوض"],
    "parcel_number": ["رقم القطعة", "رقم الأرض"],
    "feddan": ["فدان"],
    "qirat": ["قيراط", "قراط"],
    "sahm": ["سهم"],
    "land_use": ["نوع الاستخدام"],
    "crop_type": ["نوع المحصول"],
    "growth_stage": ["مراحل النمو", "مرحلة النمو"],
    "notes": ["الملاحظات"],
    "area": ["المساحة"],
}
REQUIRED = {
    "id",
    "holding_number",
    "parcel_count",
    "holder_name",
    "owner_name",
    "national_id",
    "basin_name",
    "basin_code",
    "parcel_number",
    "feddan",
    "qirat",
    "sahm",
    "land_use",
}
OPTIONAL = {
    "crop_type",
    "growth_stage",
    "notes",
    "association_name",
    "association_code",
}


class TextExtractor(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str):
        self.parts.append(data)

    def handle_starttag(self, tag, attrs):
        if tag.lower() in {"br", "p", "div", "tr", "li", "td", "th"}:
            self.parts.append("\n")

    def text(self) -> str:
        return (
            " ".join(" ".join(self.parts).split())
            if not any("\n" in p for p in self.parts)
            else "".join(self.parts)
        )


def clean_description(value: str | None) -> str:
    raw = unescape(value or "").replace("\xa0", " ")
    parser = TextExtractor()
    try:
        parser.feed(raw)
        raw = parser.text()
    except Exception:
        pass
    return re.sub(r"[ \t]+", " ", raw).strip()


def _label_pattern() -> re.Pattern:
    all_labels = sorted(
        {label for values in LABELS.values() for label in values}, key=len, reverse=True
    )
    return re.compile(
        r"(?P<label>" + "|".join(re.escape(x) for x in all_labels) + r")\s*[:：]\s*",
        re.I,
    )


LABEL_PATTERN = _label_pattern()
LABEL_TO_KEY = {
    label.casefold(): key for key, labels in LABELS.items() for label in labels
}


def _tidy(value: str) -> str:
    return value.strip(" \t\n\r|;؛،,")


def parse_description(
    description: str | None, extended: dict[str, str] | None = None
) -> dict[str, str]:
    text = clean_description(description)
    if extended:
        text += "\n" + "\n".join(
            f"{key}: {value}" for key, value in extended.items() if value
        )
    result = {key: "" for key, _ in FIELDS}
    matches = list(LABEL_PATTERN.finditer(text))
    for i, match in enumerate(matches):
        key = LABEL_TO_KEY[match.group("label").casefold()]
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        value = _tidy(text[match.end() : end])
        if key == "area":
            area_matches = list(
                re.finditer(
                    r"(فدان|قيراط|قراط|سهم)\s*[:：]?\s*([0-9٠-٩.,]+)", value, re.I
                )
            )
            for area in area_matches:
                area_key = {
                    "فدان": "feddan",
                    "قيراط": "qirat",
                    "قراط": "qirat",
                    "سهم": "sahm",
                }[area.group(1)]
                result[area_key] = _tidy(area.group(2))
        elif key in result and value:
            result[key] = value
    if not result["id"]:
        found = ANY_UUID_RE.search(text)
        if found:
            result["id"] = found.group(0)
    result["id"] = result["id"].lower().strip()
    # Area can be one phrase ("0 فدان، 4 قيراط، 0 سهم") or separate fields.
    # Resolve it from the complete text after generic parsing so Arabic commas are safe.
    for key, tokens in {"feddan": ["فدان"], "qirat": ["قيراط", "قراط"], "sahm": ["سهم"]}.items():
        for token in tokens:
            number = re.search(rf"([0-9٠-٩]+(?:[.,][0-9٠-٩]+)?)\s*{re.escape(token)}\b", text)
            if not number:
                number = re.search(rf"{re.escape(token)}\s*[:：]\s*([0-9٠-٩]+(?:[.,][0-9٠-٩]+)?)", text)
            if number:
                result[key] = number.group(1)
                break
    return result


def is_uuid(value: str) -> bool:
    return bool(UUID_RE.fullmatch((value or "").strip()))


def title_id(title: str) -> str:
    return title.strip().lower() if is_uuid(title) else ""


def element_text(element: ET.Element, tag: str) -> str:
    node = element.find(f"k:{tag}", NS)
    return (node.text or "").strip() if node is not None else ""


def set_element_text(element: ET.Element, tag: str, value: str) -> None:
    node = element.find(f"k:{tag}", NS)
    if node is None:
        node = ET.SubElement(element, f"{{{KML_NS}}}{tag}")
    node.text = value


def extract_extended_data(placemark: ET.Element) -> dict[str, str]:
    values: dict[str, str] = {}
    for data in placemark.findall(".//k:ExtendedData/k:Data", NS):
        name = data.get("name", "")
        value = data.find("k:value", NS)
        if name and value is not None and value.text:
            values[name] = value.text
    for data in placemark.findall(".//k:ExtendedData/k:SchemaData/k:SimpleData", NS):
        if data.get("name") and data.text:
            values[data.get("name", "")] = data.text
    return values


def geometry_coordinates(placemark: ET.Element) -> str:
    coords = [
        " ".join((node.text or "").split())
        for node in placemark.findall(".//k:Polygon//k:coordinates", NS)
    ]
    return " | ".join(c for c in coords if c)


def has_polygon(placemark: ET.Element) -> bool:
    return placemark.find(".//k:Polygon", NS) is not None


def canonical_description(data: dict[str, str]) -> str:
    labels = dict(FIELDS)
    lines = []
    for key, _ in FIELDS:
        value = data.get(key, "").strip()
        if value:
            lines.append(f"{labels[key]}: {value}")
    return "\n".join(lines)


@dataclass
class Parcel:
    index: int
    element: ET.Element
    title: str
    description: str
    data: dict[str, str]
    coordinates: str
    status: str = ""
    reason: str = ""
    missing: list[str] = field(default_factory=list)
    review_required: bool = False
    review_confirmed: bool = False
    retained_review: bool = False
    auto_fixed: bool = False

    @property
    def id(self) -> str:
        return self.data.get("id", "")

    @property
    def blocking(self) -> bool:
        return self.status in {
            "empty",
            "no_description",
            "incomplete",
            "unreadable",
        } or (self.review_required and not self.review_confirmed)


STATUS_LABELS = {
    "valid": "سليمة",
    "title_mismatch": "عنوان غير مطابق",
    "empty": "فارغة",
    "no_description": "بلا وصف",
    "incomplete": "بيانات ناقصة",
    "unreadable": "غير قابلة للقراءة",
    "duplicate_different": "ID مكرر وبيانات مختلفة",
    "duplicate_identical": "ID مكرر وبيانات متطابقة",
}


class KmlDocument:
    def __init__(
        self,
        source: Path,
        root: ET.Element,
        parcels: list[Parcel],
        archive: dict[str, bytes] | None = None,
        kml_member: str | None = None,
    ):
        self.source, self.root, self.parcels = source, root, parcels
        self.archive, self.kml_member = archive, kml_member

    @classmethod
    def load(cls, path: str | Path) -> "KmlDocument":
        source = Path(path)
        archive = None
        member = None
        if source.suffix.lower() == ".kmz":
            with zipfile.ZipFile(source) as zf:
                archive = {name: zf.read(name) for name in zf.namelist()}
            candidates = [name for name in archive if name.lower().endswith(".kml")]
            if not candidates:
                raise ValueError("ملف KMZ لا يحتوي على ملف KML.")
            member = next(
                (x for x in candidates if Path(x).name.lower() == "doc.kml"),
                candidates[0],
            )
            raw = archive[member]
        else:
            raw = source.read_bytes()
        try:
            root = ET.fromstring(raw)
        except ET.ParseError as exc:
            raise ValueError(f"ملف KML غير صالح: {exc}") from exc
        parcels = []
        for index, mark in enumerate(root.findall(".//k:Placemark", NS), start=1):
            if not has_polygon(mark):
                continue
            title = element_text(mark, "name")
            description = element_text(mark, "description")
            data = parse_description(description, extract_extended_data(mark))
            parcels.append(
                Parcel(
                    index, mark, title, description, data, geometry_coordinates(mark)
                )
            )
        document = cls(source, root, parcels, archive, member)
        document.analyze()
        return document

    def analyze(self) -> None:
        for parcel in self.parcels:
            retain_review = parcel.retained_review and not parcel.review_confirmed
            parcel.missing = [
                dict(FIELDS)[key]
                for key in REQUIRED
                if not parcel.data.get(key, "").strip()
            ]
            parcel.review_required = False
            if not parcel.title.strip() and not parcel.description.strip():
                parcel.status, parcel.reason = "empty", "لا يوجد عنوان أو وصف للقطعة."
            elif not parcel.description.strip():
                parcel.status, parcel.reason = (
                    "no_description",
                    "يوجد عنوان لكن لا توجد بيانات وصف.",
                )
            elif not parcel.id:
                parcel.status, parcel.reason = "incomplete", "لا يوجد ID صالح في الوصف."
            elif not is_uuid(parcel.id):
                parcel.status, parcel.reason = "unreadable", "ID الوصف ليس UUID صالحًا."
            elif parcel.missing:
                parcel.status, parcel.reason = (
                    "incomplete",
                    "حقول إلزامية ناقصة: " + "، ".join(parcel.missing),
                )
            elif title_id(parcel.title) != parcel.id:
                parcel.status, parcel.reason = (
                    "title_mismatch",
                    "عنوان القطعة لا يطابق ID الموجود في الوصف.",
                )
            else:
                parcel.status, parcel.reason = "valid", "البيانات والعنوان متطابقان."
            if retain_review and parcel.status == "valid":
                parcel.status, parcel.reason = (
                    "duplicate_identical",
                    "بيانات مكررة سابقًا؛ يلزم اعتماد المراجعة اليدوية.",
                )
                parcel.review_required = True

        groups: dict[str, list[Parcel]] = defaultdict(list)
        for parcel in self.parcels:
            if is_uuid(parcel.id) and not parcel.missing:
                groups[parcel.id].append(parcel)
        for _, group in groups.items():
            if len(group) < 2:
                continue
            fingerprint = lambda p: tuple(
                (key, p.data.get(key, "").strip()) for key, _ in FIELDS if key != "id"
            )
            identical = len({fingerprint(p) for p in group}) == 1
            for parcel in group[1:]:
                parcel.status = (
                    "duplicate_identical" if identical else "duplicate_different"
                )
                parcel.reason = "ID مكرر وبيانات الوصف " + (
                    "متطابقة؛ يلزم اعتماد المراجعة."
                    if identical
                    else "مختلفة؛ سيُنشأ ID جديد."
                )
                parcel.review_required = identical
                parcel.retained_review = identical

    def apply_safe_fixes(self) -> int:
        changed = 0
        groups: dict[str, list[Parcel]] = defaultdict(list)
        for parcel in self.parcels:
            if is_uuid(parcel.id) and not parcel.missing:
                groups[parcel.id].append(parcel)
        for _, group in groups.items():
            fingerprint = lambda p: tuple(
                (key, p.data.get(key, "").strip()) for key, _ in FIELDS if key != "id"
            )
            identical = len(group) > 1 and len({fingerprint(p) for p in group}) == 1
            for parcel in group[1:]:
                if identical:
                    parcel.retained_review = True
                self._set_id(parcel, str(uuid.uuid4()))
                parcel.auto_fixed = True
                changed += 1
        for parcel in self.parcels:
            if (
                is_uuid(parcel.id)
                and not parcel.missing
                and title_id(parcel.title) != parcel.id
            ):
                parcel.title = parcel.id
                set_element_text(parcel.element, "name", parcel.id)
                parcel.auto_fixed = True
                changed += 1
        self.analyze()
        # Duplicate identical records retain an explicit review requirement after IDs are made unique.
        for parcel in self.parcels:
            if parcel.auto_fixed and parcel.status == "valid":
                parcel.reason = "تم إصلاحها تلقائيًا."
        return changed

    def _set_id(self, parcel: Parcel, value: str) -> None:
        parcel.data["id"] = value.lower()
        parcel.title = value.lower()
        parcel.description = canonical_description(parcel.data)
        set_element_text(parcel.element, "name", parcel.title)
        set_element_text(parcel.element, "description", parcel.description)

    def update_parcel(
        self, parcel: Parcel, values: dict[str, str], confirmed: bool = False
    ) -> None:
        parcel.data.update({key: values.get(key, "").strip() for key, _ in FIELDS})
        parcel.title = parcel.data.get("id", "")
        parcel.description = canonical_description(parcel.data)
        set_element_text(parcel.element, "name", parcel.title)
        set_element_text(parcel.element, "description", parcel.description)
        self.analyze()
        refreshed = self.parcels[self.parcels.index(parcel)]
        if confirmed and not refreshed.missing and is_uuid(refreshed.id):
            refreshed.retained_review = False
            refreshed.review_required = False
            refreshed.review_confirmed = True
            refreshed.status, refreshed.reason = "valid", "تم اعتماد المراجعة اليدوية."

    def apply_association(
        self, name: str, code: str, selected: list[Parcel] | None = None
    ) -> None:
        targets = selected if selected is not None else self.parcels
        for parcel in targets:
            if name.strip():
                parcel.data["association_name"] = name.strip()
            if code.strip():
                parcel.data["association_code"] = code.strip()
            parcel.description = canonical_description(parcel.data)
            set_element_text(parcel.element, "description", parcel.description)
        self.analyze()

    def summary(self) -> Counter:
        counts = Counter(parcel.status for parcel in self.parcels)
        counts["total"] = len(self.parcels)
        counts["blocking"] = sum(p.blocking for p in self.parcels)
        counts["auto_fixed"] = sum(p.auto_fixed for p in self.parcels)
        return counts

    def _style_id(self, parcel: Parcel) -> str:
        if parcel.status in {"empty", "no_description", "incomplete", "unreadable"}:
            return "parcel-red"
        if parcel.review_required and not parcel.review_confirmed:
            return "parcel-yellow"
        assoc = parcel.data.get("association_name", "")
        if "إصلاح" in assoc or "اصلاح" in assoc:
            return "parcel-cyan"
        if "ائتمان" in assoc:
            return "parcel-blue"
        return "parcel-blue"

    def _install_styles(self) -> None:
        container = self.root.find("k:Document", NS) or self.root
        for old in list(container.findall("k:Style", NS)):
            if old.get("id", "").startswith("parcel-"):
                container.remove(old)
        styles = {
            "parcel-blue": ("ccff6600", "ffcc6600"),
            "parcel-cyan": ("ccffff00", "ffffff00"),
            "parcel-yellow": ("cc00ffff", "ff00ffff"),
            "parcel-red": ("ff0000ff", "ff0000ff"),
        }
        for style_id, (fill, line) in styles.items():
            style = ET.Element(f"{{{KML_NS}}}Style", {"id": style_id})
            poly = ET.SubElement(style, f"{{{KML_NS}}}PolyStyle")
            ET.SubElement(poly, f"{{{KML_NS}}}color").text = fill
            ET.SubElement(poly, f"{{{KML_NS}}}fill").text = "1"
            ET.SubElement(poly, f"{{{KML_NS}}}outline").text = "1"
            line_style = ET.SubElement(style, f"{{{KML_NS}}}LineStyle")
            ET.SubElement(line_style, f"{{{KML_NS}}}color").text = line
            ET.SubElement(line_style, f"{{{KML_NS}}}width").text = "1"
            container.insert(0, style)

    def export_kml(self, output: str | Path) -> None:
        self._install_styles()
        for parcel in self.parcels:
            style = parcel.element.find("k:styleUrl", NS)
            if style is None:
                style = ET.SubElement(parcel.element, f"{{{KML_NS}}}styleUrl")
            style.text = "#" + self._style_id(parcel)
            set_element_text(
                parcel.element, "name", parcel.data.get("id", parcel.title)
            )
            set_element_text(
                parcel.element, "description", canonical_description(parcel.data)
            )
        raw = ET.tostring(self.root, encoding="utf-8", xml_declaration=True)
        output = Path(output)
        if output.suffix.lower() == ".kmz":
            archive = copy.copy(self.archive) if self.archive else {}
            archive[self.kml_member or "doc.kml"] = raw
            with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as zf:
                for name, data in archive.items():
                    zf.writestr(name, data)
        else:
            output.write_bytes(raw)

    def export_excel(self, output: str | Path) -> None:
        if any(parcel.blocking for parcel in self.parcels):
            raise ValueError("لا يمكن بناء Excel قبل معالجة كل القطع المانعة.")
        book = Workbook()
        sheet = book.active
        sheet.title = "القطع"
        sheet.sheet_view.rightToLeft = True
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for col, (_, label) in enumerate(FIELDS, start=1):
            cell = sheet.cell(1, col, label)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for row, parcel in enumerate(self.parcels, start=2):
            for col, (key, _) in enumerate(FIELDS, start=1):
                sheet.cell(row, col, parcel.data.get(key, ""))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for col in range(1, len(FIELDS) + 1):
            letter = get_column_letter(col)
            sheet.column_dimensions[letter].width = min(
                38,
                max(
                    14,
                    max(
                        len(str(sheet.cell(r, col).value or ""))
                        for r in range(1, sheet.max_row + 1)
                    )
                    + 2,
                ),
            )
        book.save(output)
